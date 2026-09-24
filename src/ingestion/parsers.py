"""Parsers for supported document formats."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Iterable

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from src.config import SUPPORTED_EXTENSIONS
from src.models import DocumentSection


class DocumentParseError(ValueError):
    """Raised when an uploaded document cannot be safely parsed."""


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("The text encoding could not be determined.")


def _batch_lines(
    source: str,
    rows: Iterable[str],
    *,
    batch_size: int,
    location_prefix: str,
) -> list[DocumentSection]:
    output: list[DocumentSection] = []
    batch: list[str] = []
    start_row = 1
    current_row = 0

    for current_row, row in enumerate(rows, start=1):
        batch.append(row)
        if len(batch) >= batch_size:
            output.append(
                DocumentSection(
                    source=source,
                    location=f"{location_prefix} {start_row}-{current_row}",
                    text="\n".join(batch),
                )
            )
            batch = []
            start_row = current_row + 1

    if batch:
        output.append(
            DocumentSection(
                source=source,
                location=f"{location_prefix} {start_row}-{current_row}",
                text="\n".join(batch),
            )
        )
    return output


def _parse_pdf(name: str, data: bytes) -> list[DocumentSection]:
    try:
        reader = PdfReader(io.BytesIO(data))
        sections = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                sections.append(
                    DocumentSection(source=name, location=f"page {index}", text=text)
                )
        if not sections:
            raise DocumentParseError(
                "No selectable text was found in the PDF. Scanned PDFs need OCR."
            )
        return sections
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"Could not read PDF '{name}'.") from exc


def _parse_docx(name: str, data: bytes) -> list[DocumentSection]:
    try:
        document = Document(io.BytesIO(data))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        if not paragraphs:
            raise DocumentParseError("No readable paragraphs were found in the DOCX file.")
        return _batch_lines(
            name,
            paragraphs,
            batch_size=25,
            location_prefix="paragraphs",
        )
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"Could not read DOCX '{name}'.") from exc


def _parse_txt(name: str, data: bytes) -> list[DocumentSection]:
    text = _decode_text(data).strip()
    if not text:
        raise DocumentParseError("The text file is empty.")
    return [DocumentSection(source=name, location="document", text=text)]


def _parse_csv(name: str, data: bytes) -> list[DocumentSection]:
    text = _decode_text(data)
    try:
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    except csv.Error as exc:
        raise DocumentParseError(f"Could not read CSV '{name}'.") from exc

    if not rows:
        raise DocumentParseError("The CSV file is empty.")

    header = rows[0]
    rendered_rows = []
    for row_number, row in enumerate(rows[1:], start=2):
        pairs = []
        for index, value in enumerate(row):
            key = header[index] if index < len(header) and header[index] else f"column_{index + 1}"
            pairs.append(f"{key}: {value}")
        rendered_rows.append(f"Row {row_number} | " + " | ".join(pairs))

    if not rendered_rows:
        rendered_rows = ["Columns: " + ", ".join(header)]

    return _batch_lines(
        name,
        rendered_rows,
        batch_size=40,
        location_prefix="rows",
    )


def _parse_xlsx(name: str, data: bytes) -> list[DocumentSection]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections: list[DocumentSection] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(value) if value is not None else "" for value in rows[0]]
            rendered = []
            for row_number, row in enumerate(rows[1:], start=2):
                values = ["" if value is None else str(value) for value in row]
                if not any(values):
                    continue
                pairs = []
                for index, value in enumerate(values):
                    key = header[index] if index < len(header) and header[index] else f"column_{index + 1}"
                    pairs.append(f"{key}: {value}")
                rendered.append(f"Row {row_number} | " + " | ".join(pairs))

            sheet_sections = _batch_lines(
                name,
                rendered or ["Columns: " + ", ".join(header)],
                batch_size=35,
                location_prefix=f"sheet '{sheet.title}' rows",
            )
            sections.extend(sheet_sections)

        if not sections:
            raise DocumentParseError("No readable spreadsheet data was found.")
        return sections
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"Could not read XLSX '{name}'.") from exc


def parse_document(name: str, data: bytes) -> list[DocumentSection]:
    """Parse raw uploaded bytes into source-aligned sections."""

    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(
            f"Unsupported file type '{suffix or 'unknown'}'. Supported: PDF, DOCX, TXT, CSV, XLSX."
        )

    parser = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".txt": _parse_txt,
        ".csv": _parse_csv,
        ".xlsx": _parse_xlsx,
    }[suffix]
    return parser(name, data)
